from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor

path = Path(__file__).resolve().parents[1] / 'BPM EA and AISD .docx'
doc = Document(path)

for paragraph in doc.paragraphs:
    if paragraph.text.strip() == '3 Justification of Decisions':
        paragraph.text = '3 Justification of Decisions Summary'
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
        break

def add_heading(text, level):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph

def add_body(text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(7)
    for run in paragraph.runs:
        run.font.size = Pt(10)
    return paragraph

add_heading('4 Detailed Justification of Decisions', 1)
add_body('This section explains the design choices represented in the executable BPMN files. The justification refers to Process_Hospital_Master and the five linked processes: Process_Referral, Process_Appointment, Process_Finance, Process_Treatment and Process_EnquiryLetter.')

add_heading('4.1 Process Structure', 2)
add_body('The solution uses one master process and five smaller executable child processes. Process_Hospital_Master starts the patient pathway and calls Referral and Consultant Review first. Only an ACCEPT referral decision permits the Appointment and Patient Communication process to begin. The master process then coordinates Funding Payment Refund and Treatment Chemotherapy before calling the Clinic Letter and Pathway Monitoring route. This structure preserves an end-to-end patient journey without forcing all activities, exception paths and message flows into one unreadable diagram.')
add_body('The Enquiry and Clinic Letter process contains separate start events because patient enquiries can arise independently, while Clinic Letter production begins after treatment or consultation information is available. This is intentional: an enquiry must be handled without waiting for the full treatment path, whereas the Clinic Letter path depends on clinical information and approval. The master process therefore shows the planned care pathway, while the child processes retain the operational detail needed by each team.')

add_heading('4.2 Participant Boundaries', 2)
add_body('Participant boundaries follow authority rather than convenience. Referral and Consultant Review separates the Referring Organisation from the hospital and uses Medical Secretary and Consultant Clinical Team lanes. The Medical Secretary records information and manages completeness, but the Consultant makes ACCEPT, REJECT, MORE INFORMATION and REDIRECT decisions. This prevents an administrative role from making a clinical acceptance decision.')
add_body('Appointment activities are assigned to outpatient booking and patient communication roles. Treatment and chemotherapy activities remain within the clinical and treatment-booking boundary. Finance owns funding route selection, payment investigation, insurer approval and refund authority. Enquiry handling routes administrative, financial and clinical enquiries to different responsible teams. These boundaries make responsibility visible in the BPMN and support auditability when a task, decision or exception must be explained.')

add_heading('4.3 Task Allocation', 2)
add_body('User tasks are used where a person must apply judgement, record consent, approve a decision, contact a patient, confirm a booking or manage an exception. Examples include Perform Consultant Review, Confirm Appointment, Record Patient Consent, Review Refund Authority and Approve Clinical Content. The associated Camunda Forms collect the variables needed for the next gateway or external worker, while keeping accountable decisions with a named human role.')
add_body('Service tasks are used only for repeatable simulated external services. The Referral process now invokes Validate Referral Reference with External Intake Service using job type validate-referral-reference. Appointment invokes search-appointment-slot and send-appointment-letter. Treatment invokes check-treatment-resource. Finance invokes process-payment and process-refund. Clinic Letter distribution invokes send-clinic-letter. This allocation demonstrates an executable Camunda design: workers return variables such as referralReferenceValidated, slotAvailable, treatmentAvailable, paymentStatus and clinicLetterSent for the following BPMN decision.')

add_heading('4.4 Gateways and Business Rules', 2)
add_body('Exclusive gateways make business rules explicit and testable. In Referral, Supporting Information Complete prevents consultant review until informationComplete is true, and Referral Decision routes ACCEPT, REJECT, MORE_INFORMATION and REDIRECT outcomes. In Appointment, Suitable Slot Available uses slotAvailable returned by the scheduling worker. If a slot exists, Appointment Within Two Weeks uses appointmentWithinTwoWeeks to decide whether telephone contact is required. Telephone Outcome routes ANSWERED, NO_ANSWER, WRONG_NUMBER and RESCHEDULE outcomes, with rescheduling returning to the slot search.')
add_body('Treatment uses authorisation, consent, resource-availability, clinical-decision and financial-impact gateways. Finance selects the hospital, insurer or patient funding route and treats CONFIRMATION_LOST as an investigation path rather than an automatic repeat charge. The refund worker is reachable only after refundRequired and refundApproved conditions are satisfied. Enquiry and Clinic Letter uses enquiry type, urgent clinical concern, approval status, seven-day readiness and age-band gateways. These rules prevent unsafe shortcuts, such as delivering treatment without consent, distributing an unapproved letter or refunding without Finance approval.')

add_heading('4.5 External Interactions', 2)
add_body('Message flows distinguish communication with external organisations from internal hospital control flow. The Referral model communicates with the Referring Organisation for the referral, missing-information request and additional information. Appointment represents the Scheduling Service, Correspondence Service and patient communication. Finance represents insurer and payment-provider interactions. Treatment represents treatment, laboratory or imaging resource services. Clinic Letter distribution represents correspondence recipients and the monitoring route. The diagram therefore identifies the system boundary and makes clear which dependency is external.')
add_body('Java external workers are deliberately simulated for coursework. They do not claim to connect to real hospital systems. Instead, each worker reads controlled input variables from the Form, writes an auditable output and lets the BPMN gateway choose the next path. This allows the group to demonstrate worker execution, variable propagation and error-oriented process design without handling real patient data or creating unsafe real integrations.')

add_heading('4.6 Exception Handling', 2)
add_body('The process does not treat every problem as a simple end event. Missing referral information returns to a request and receipt loop. A no-slot appointment is recorded, reviewed for urgency and can return to a revised search. Telephone failures are recorded separately so that an answered contact is not confused with a wrong number or a reschedule request. Treatment resource unavailability enters a pending and notification route with a retry task rather than silently continuing.')
add_body('Finance distinguishes declined, cancelled and incomplete payments from confirmation loss. Confirmation loss leads to investigation because automatically charging again could create a duplicate payment. Refunds require a human authority check. In the Clinic Letter path, unapproved or delayed letters enter correction, weekly reminder, monitoring and escalation routes, including the one-month and three-month escalation levels. Urgent clinical enquiries are immediately escalated to an authorised clinical role. These paths maintain visibility and control when the normal pathway cannot continue.')

add_heading('4.7 Assumptions', 2)
add_body('The model assumes that the Camunda Forms are deployed before their BPMN processes and that the master process calls the latest deployment of each child process. It assumes that external service outcomes are simulated through form variables and Java workers, that requested appointment dates can be compared with the current date, and that an appointment within fourteen days requires telephone contact. It also assumes that clinical staff retain authority for clinical decisions, Finance retains authority for refunds, and a Clinic Letter cannot be distributed until clinical approval has been recorded.')
add_body('The model also assumes that a referral reference can be validated by a simulated external intake service and that a false validation result is recorded for follow-up rather than automatically rejecting a patient. These assumptions are documented so that a future production implementation can replace simulated services with authenticated integrations and agreed operational policies.')

add_heading('4.8 Alternatives Considered', 2)
add_body('One alternative was a single large BPMN model containing referral, appointment, finance, treatment, enquiries and clinic letters. This would keep every activity on one canvas, but it would create excessive crossings, reduce readability and make it difficult for each team member to explain their own executable responsibility. The team instead selected five independent executable models with a master process for orchestration.')
add_body('A second alternative was to model every external interaction as a manual task. That approach would be simpler to draw but would fail the requirement to demonstrate External Workers and would hide the difference between a human decision and a technical service. A third alternative was automatic retry for failed payment or booking operations. The chosen design requires a user task to revise the input or investigate the result because the health and financial contexts make uncontrolled repetition inappropriate.')

add_heading('4.9 Trade Offs', 2)
add_body('Splitting the solution into child processes improves readability, team ownership, independent deployment and demonstration. The trade-off is integration overhead: process IDs, Call Activity bindings, Form IDs, worker job types and variable names must remain consistent. The repository addresses this by keeping editable BPMN and Forms together, centralising worker registration in HospitalExternalWorkers and documenting deployment order and test cases.')
add_body('Using simulated workers improves reproducibility and avoids real patient or payment data, but it does not prove production interoperability. Using explicit exception tasks improves safety and auditability, but it adds steps and makes diagrams larger. These trade-offs are acceptable for the coursework release because the objective is an executable, understandable and testable operational BPMN baseline rather than a live hospital deployment.')

try:
    doc.save(path)
except PermissionError:
    doc.save(path.with_name('BPM EA and AISD Detailed Justification.docx'))
