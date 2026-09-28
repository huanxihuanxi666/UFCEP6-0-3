from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'D:/Camunda-Workspace/models/assignment')
OUT=ROOT/'output'/'Hospital总图与五人子图逻辑及Java代码说明.docx'
J=ROOT/'External_Workers/src/main/java/org/hospital/workers'

doc=Document(); section=doc.sections[0]
section.top_margin=Cm(1.65); section.bottom_margin=Cm(1.65); section.left_margin=Cm(1.65); section.right_margin=Cm(1.65)

def font(style, size, bold=False, face='Microsoft YaHei'):
    style.font.name=face; style.font.size=Pt(size); style.font.bold=bold; style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr(); rf=rpr.rFonts
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for k in ('ascii','hAnsi','eastAsia'): rf.set(qn('w:'+k),face)
font(doc.styles['Normal'],10); font(doc.styles['Title'],19,True); font(doc.styles['Heading 1'],14,True); font(doc.styles['Heading 2'],11,True)
doc.styles['Normal'].paragraph_format.space_after=Pt(4); doc.styles['Normal'].paragraph_format.line_spacing=1.12
code_style=doc.styles.add_style('HospitalCode', WD_STYLE_TYPE.PARAGRAPH)
font(code_style,7.2,False,'Consolas'); code_style.paragraph_format.space_after=Pt(0); code_style.paragraph_format.line_spacing=1.0

def shade(cell,color):
    p=cell._tc.get_or_add_tcPr(); el=OxmlElement('w:shd'); el.set(qn('w:fill'),color); p.append(el)
def borders(cell):
    tc=cell._tc.get_or_add_tcPr(); bd=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); bd.append(e)
    tc.append(bd)
def matrix(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(headers): t.rows[0].cells[i].text=x
    for vals in rows:
        cells=t.add_row().cells
        for i,x in enumerate(vals): cells[i].text=x
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(cell,'173E48' if ri==0 else ('EEF4F3' if ri%2==0 else 'FFFFFF')); borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.05
                for run in p.runs:
                    run.font.name='Microsoft YaHei'; run.font.size=Pt(8.5); run.bold=(ri==0); run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    return t
def code(text):
    p=doc.add_paragraph(style='HospitalCode'); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(5)
    p.add_run(text)
    ppr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F3F5F7'); ppr.append(sh)
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)
def clean_title(p):
    ppr=p._p.get_or_add_pPr(); bd=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom'); b.set(qn('w:val'),'nil'); bd.append(b); ppr.append(bd)

title=doc.add_paragraph('Hospital BPMN Logic and Java Code Reference',style='Title'); clean_title(title)
doc.add_paragraph('用途：答辩时解释总图如何调用五张子图、每位成员图中的判断条件，以及哪个 Java 文件让自动节点继续运行。本文按当前 BPMN 与 Java 源码编写。')
doc.add_heading('阅读顺序',1)
matrix(['先看什么','答案'],[
 ('主流程入口','只启动 Process_Hospital_Master。它依次调用转诊、预约、资金、治疗和信件子流程。'),
 ('五位成员','邓金旺：转诊；吴家威：预约；吴同宇：资金；林佳强：治疗；肖宇莹：咨询和信件。'),
 ('Java 启动程序','HospitalExternalWorkers.java 是唯一 main 程序；它一次注册全部六种 job type。'),
 ('变量传递','五个 Call Activity 都使用 propagateAllChildVariables=true，因此子图判断结果会回到总图。'),
], [3.4,13.0])
doc.add_paragraph('结论：总图负责跨模块顺序和三种结束结果；每张子图负责本模块的详细分支；Java worker 只处理 BPMN Service Task，不代替 Tasklist 中的人工作业。')

page('1 Total Diagram Logic')
doc.add_paragraph('总图文件：bpmn/Hospital_Master_Process.bpmn。流程 ID：Process_Hospital_Master。总图只有一个启动点和三个结束点。')
matrix(['次序','总图节点','依据变量与结果'],[
 ('1','Referral and Consultant Review','调用 Process_Referral。referralDecision=ACCEPT 才进入预约；其他值直接到 Referral Pathway Closed。'),
 ('2','Appointment and Patient Communication','调用 Process_Appointment。预约子图结束后才进入资金路径。'),
 ('3','Funding Payment Refund','调用 Process_Finance。资金清算网关检查 fundingType、fundingApproved、paymentStatus、urgentTreatmentOverride。'),
 ('4','Treatment and Chemotherapy','资金通过才调用 Process_Treatment。治疗子图结束后进入信件模块。'),
 ('5','Enquiry Clinic Letter Monitoring','调用 Process_EnquiryLetter。信件子图结束后到 Patient Pathway Completed。'),
], [1.0,5.5,9.9])
doc.add_heading('三个结束点',2)
matrix(['结束事件','到达条件'],[
 ('Referral Pathway Closed','referralDecision 不是 ACCEPT。'),
 ('Pathway Closed Pending Follow up','资金清算条件不成立。例如 fundingType=INSURER 且 fundingApproved=false。'),
 ('Patient Pathway Completed','转诊接受、资金清算、治疗完成、信件子图完成。'),
], [5.6,10.8])
doc.add_heading('总图的关键 FEEL 条件',2)
code('Referral accepted:  referralDecision = "ACCEPT"\n\nFinancially cleared:  fundingType = "HOSPITAL"\n  or fundingType = "EXEMPTION"\n  or (fundingType = "INSURER" and fundingApproved = true)\n  or (fundingType = "PATIENT" and (paymentStatus = "SUCCESSFUL" or urgentTreatmentOverride = true))')

page('2 Deng Jinwang Referral and Consultant Review')
doc.add_paragraph('子图文件：bpmn/23084366_Referral_Consultant_Review/BPMN/Referral_Consultant_Review.bpmn。流程 ID：Process_Referral。这个模块没有 Service Task，因此 ReferralWorkers.java 只明确说明“不注册 worker”。')
matrix(['节点','判断变量','走向'],[
 ('Record Referral and Supporting Documents','informationComplete、urgentReferral','资料进入检查；紧急信息会触发临床升级路径。'),
 ('Check Supporting Information','资料是否完整','不完整时回到资料补充。'),
 ('Perform Consultant Review','referralDecision','REJECT：转诊关闭；MORE_INFORMATION：补资料后再审；ACCEPT：返回总图预约；REDIRECT：转往其他路径。'),
], [5.5,4.0,6.9])
doc.add_paragraph('与总图的交接：最关键输出是 referralDecision。总图不读取任何转诊任务名称，只读取这个变量。')

page('3 Wu Jiawei Appointment and Patient Communication')
doc.add_paragraph('子图文件：bpmn/23084364_Appointment_Patient_Communication/BPMN/Appointment_Patient_Communication.bpmn。流程 ID：Process_Appointment。')
matrix(['节点','变量或自动任务','逻辑'],[
 ('Receive Booking Request','requestedSlotAvailable','提交预约信息后，进入 search-appointment-slot。'),
 ('Search External Scheduling Service','Java job: search-appointment-slot','AppointmentWorkers 将 requestedSlotAvailable 转成 slotAvailable，并计算 appointmentWithinTwoWeeks。'),
 ('Suitable Slot Available?','slotAvailable','true：确认预约；false：记录无合适时段、升级或重新搜索。'),
 ('Appointment Within Two Weeks?','appointmentWithinTwoWeeks','true：直接发送预约信；false：电话联系病人。'),
 ('Telephone Outcome','contactOutcome','ANSWERED 完成沟通；RESCHEDULE 回到搜号源；WRONG_NUMBER 记录错误号码。'),
 ('Send Appointment Letter','Java job: send-appointment-letter','worker 写入 appointmentLetterSent=true；子图完成，回到总图资金模块。'),
], [5.2,4.8,6.4])

page('4 Wu Tongyu Funding Payment and Refund')
doc.add_paragraph('子图文件：bpmn/23084362_Funding_Payment_Refund/BPMN/Funding_Payment_Refund.bpmn。流程 ID：Process_Finance。')
matrix(['节点','变量或自动任务','逻辑'],[
 ('Funding Type','fundingType','HOSPITAL / EXEMPTION：直接清算；INSURER：等待批准；PATIENT：调用支付服务。'),
 ('Insurer Funding Approved?','fundingApproved','true：清算；false：Funding Pending，回到总图后进入 Pending Follow up 结束。'),
 ('Send Secure Payment Request','Java job: process-payment','FinanceWorkers 读取 paymentScenario，并写入 paymentStatus。'),
 ('Payment Status','paymentStatus','SUCCESSFUL：清算；DECLINED / CANCELLED / INCOMPLETE：重试；CONFIRMATION_LOST / DUPLICATE_RISK：调查后判断紧急覆盖。'),
 ('Urgent Treatment Despite Unconfirmed Payment?','urgentTreatmentOverride','true：记录临床覆盖；false：保持未清算。'),
 ('Refund Approved?','refundRequired、refundApproved','同时为 true 时调用 process-refund；否则退款不执行。'),
], [5.2,4.8,6.4])

page('5 Lin Jiaqiang Treatment and Chemotherapy')
doc.add_paragraph('子图文件：bpmn/23084368_Treatment_Chemotherapy/BPMN/Treatment_Chemotherapy.bpmn。流程 ID：Process_Treatment。')
matrix(['节点','变量或自动任务','逻辑'],[
 ('Clinical Assessment','patientConsent','true：记录同意并创建治疗请求；false：治疗不继续。'),
 ('Treatment Request','treatmentAuthorised','true：检查资源；false：退回并重新授权。'),
 ('Check Treatment Resources','Java job: check-treatment-resource','TreatmentWorkers 将 requestedTreatmentAvailable 写成 treatmentAvailable。'),
 ('Resources Available?','treatmentAvailable','true：预约治疗；false：保持待处理、更新资源后重试。'),
 ('Treatment Cycle Review','moreCycles','true：进入血检和临床复审；false：Treatment Complete Clinic Letter Handover。'),
 ('Clinical Decision','clinicalDecision','CONTINUE：下一周期；DELAY：延迟后复审；MODIFY：修改治疗并判断 financialImpact。'),
], [5.2,4.8,6.4])
doc.add_paragraph('与总图的交接：治疗子图无论“完成治疗”或“临床不继续”都能结束，然后总图调用信件模块；总图没有额外读取治疗变量。')

page('6 Xiao Yuying Enquiry Clinic Letter Monitoring')
doc.add_paragraph('子图文件：bpmn/24000813_Enquiry_ClinicLetter_Monitoring/BPMN/Enquiry_ClinicLetter_Monitoring.bpmn。流程 ID：Process_EnquiryLetter。')
matrix(['节点','变量或自动任务','逻辑'],[
 ('Enquiry Type','enquiryType','ADMINISTRATIVE、FINANCIAL、CLINICAL 分别交给相应人员；urgentClinicalConcern=true 会立即升级。'),
 ('Prepare Clinic Letter','clinicLetterWithin7Days','true：进入临床内容审批；false：加入监控并进入每周提醒。'),
 ('Clinical Content Approved?','clinicLetterApproved','true：行政核对收件人；false：重新判断是否仍可在七天内发出。'),
 ('Outstanding Duration','letterAgeBand','7D_TO_1M：提醒顾问；OVER_1M：升级行政经理；OVER_3M：升级高层。'),
 ('Clinic Letter Completed?','clinicLetterCompleted','true：回到审批；false：继续每周监控循环。'),
 ('Distribute Clinic Letter','Java job: send-clinic-letter','worker 要求 clinicLetterApproved=true；完成后写入 clinicLetterSent=true。'),
], [5.2,4.8,6.4])

page('7 Java Program and Worker Registration')
doc.add_paragraph('启动程序：External_Workers/src/main/java/org/hospital/workers/HospitalExternalWorkers.java。main 方法使用 ZEEBE_ADDRESS（默认 127.0.0.1:26500）创建客户端，然后注册所有模块 worker，并保持进程运行。')
matrix(['Java 文件','负责的 BPMN job type','关键输出'],[
 ('AppointmentWorkers.java','search-appointment-slot；send-appointment-letter','slotAvailable；appointmentWithinTwoWeeks；appointmentLetterSent'),
 ('TreatmentWorkers.java','check-treatment-resource','treatmentAvailable'),
 ('FinanceWorkers.java','process-payment；process-refund','paymentStatus；refundStatus'),
 ('EnquiryClinicLetterWorkers.java','send-clinic-letter','clinicLetterSent'),
 ('ReferralWorkers.java','无','转诊图没有 Service Task，因此不订阅 job。'),
 ('WorkerSupport.java','所有 job 共用','注册、完成 job、失败重试、变量读写工具。'),
], [4.5,6.2,5.7])
doc.add_paragraph('运行关系：Tasklist 完成人工任务 → BPMN 到达 Service Task → 对应 Java worker 接收 job → worker 写回变量 → 网关依据变量决定下一条线。')

page('Appendix A Full Java Source Files')
doc.add_paragraph('以下是当前项目中的完整 Java worker 源码。HospitalExternalWorkers.java 是要启动的主程序。')
for filename in ['HospitalExternalWorkers.java','WorkerSupport.java','ReferralWorkers.java','AppointmentWorkers.java','TreatmentWorkers.java','FinanceWorkers.java','EnquiryClinicLetterWorkers.java']:
    doc.add_heading(filename,2)
    code((J/filename).read_text(encoding='utf-8'))

OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(OUT)
