from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(r'D:/Camunda-Workspace/models/assignment/output/总图三条路径_Tasklist填写运行手册.docx')
doc=Document(); sec=doc.sections[0]
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)

def style(s,size,bold=False):
    s.font.name='Microsoft YaHei'; s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=RGBColor(0,0,0)
    r=s.element.get_or_add_rPr(); f=r.rFonts
    if f is None: f=OxmlElement('w:rFonts'); r.insert(0,f)
    for k in ('ascii','hAnsi','eastAsia'): f.set(qn('w:'+k),'Microsoft YaHei')
style(doc.styles['Normal'],10); style(doc.styles['Title'],19,True); style(doc.styles['Heading 1'],13,True)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)

def table(rows):
    heads=['成员','Tasklist 任务','只需勾选／选择']
    widths=[2.2,5.0,9.0]
    t=doc.add_table(rows=1, cols=3); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(heads): t.rows[0].cells[i].text=x
    for values in rows:
        cells=t.add_row().cells
        for i,x in enumerate(values): cells[i].text=x
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'173E48' if ri==0 else ('EEF4F3' if ri%2==0 else 'FFFFFF')); props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
                for run in p.runs:
                    run.font.name='Microsoft YaHei'; run.font.size=Pt(9); run.bold=(ri==0); run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)

title=doc.add_paragraph('总图三条路径：Form 只填选择项',style='Title')
pr=title._p.get_or_add_pPr(); borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'nil'); borders.append(bottom); pr.append(borders)
doc.add_paragraph('只启动总图。Tasklist 出现下列任务时，按这一页选择或勾选；所有文字、日期、编号、备注等自由输入项都不用看。没有列出的任务：直接 Complete Task。')

def route(name,explain,rows):
    doc.add_page_break(); doc.add_heading(name,1); doc.add_paragraph(explain); table(rows)

route('路径一｜结束：Referral Pathway Closed','最短路径。只出现邓金旺的两个 Form；顾问拒绝后总图立即结束。',[
    ('邓金旺','Referral Intake','Information complete：✓ 勾选\nUrgent referral：✗ 不勾选'),
    ('邓金旺','Consultant Review','Referral decision：选择 REJECT'),
])

route('路径二｜结束：Pathway Closed Pending Follow up','转诊和预约均通过；保险不批准时，资金闸门结束总图。',[
    ('邓金旺','Referral Intake','Information complete：✓ 勾选\nUrgent referral：✗ 不勾选'),
    ('邓金旺','Consultant Review','Referral decision：选择 ACCEPT'),
    ('吴家威','Appointment Booking','Speciality：选择 Oncology\nAppointment priority：选择 Routine\nPreferred timeframe：选择 Next available\nRequested slot available：✓ 勾选'),
    ('吴同宇','Funding Payment','Funding type：选择 INSURER\nPayment scenario：选择 SUCCESSFUL'),
    ('吴同宇','Insurer Approval','Funding approved：✗ 不勾选'),
])

route('路径三｜结束：Patient Pathway Completed','完整路径。依次由五位成员操作；全部按下列值选择即可到达总图最终结束。',[
    ('邓金旺','Referral Intake','Information complete：✓ 勾选\nUrgent referral：✗ 不勾选'),
    ('邓金旺','Consultant Review','Referral decision：选择 ACCEPT'),
    ('吴家威','Appointment Booking','Speciality：选择 Oncology\nAppointment priority：选择 Routine\nPreferred timeframe：选择 Next available\nRequested slot available：✓ 勾选'),
    ('吴同宇','Funding Payment','Funding type：选择 HOSPITAL\nPayment scenario：选择 SUCCESSFUL'),
    ('林佳强','Clinical Assessment','Patient consent：✓ 勾选'),
    ('林佳强','Treatment Request','Proposed treatment：选择 Chemotherapy\nCycle frequency：选择 Weekly\nTreatment authorised：✓ 勾选\nRequested treatment available：✓ 勾选'),
    ('林佳强','Treatment Cycle Review（若出现）','More cycles：✗ 不勾选\nClinical decision：选择 CONTINUE\nFinancial impact：✗ 不勾选'),
    ('肖宇莹','Clinic Letter','Within 7 days：✓ 勾选'),
    ('肖宇莹','Clinic Letter Approval','Clinical content approved：✓ 勾选\nWithin 7 days：✓ 勾选'),
])

doc.add_page_break(); doc.add_heading('最后只看三个结束点',1)
table([
    ('路径一','Referral Pathway Closed','看到此结束点即成功'),
    ('路径二','Pathway Closed Pending Follow up','看到此结束点即成功'),
    ('路径三','Patient Pathway Completed','看到此结束点即成功'),
])
doc.add_paragraph('说明：以上每一行仅保留会影响网关走向的下拉选项和勾选框。其他空格文本随意填写或保持默认即可。')
OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(OUT)
